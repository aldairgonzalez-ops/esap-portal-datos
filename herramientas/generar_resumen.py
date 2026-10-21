# -*- coding: utf-8 -*-
r"""
Genera el Resumen Ejecutivo de la base de datos ESAP (Word) a partir de la misma extracción del portal.

Las cifras salen de la extracción; el texto de hallazgos se arma con esas cifras y la sección final
cita el portal (con enlaces directos a cada vista) y el repositorio con el detalle técnico.

Uso:
  python herramientas/generar_resumen.py --extraccion extracciones/2026-10-21 --version v1.0 \
      --url-portal "https://…/Portal_de_Datos_ESAP_v1.0.html" --repo "https://github.com/…"
Salida: informe/Resumen_Ejecutivo_BD_ESAP_<version>.docx
"""
import argparse
import datetime
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import modelo as M  # noqa: E402

AZUL = RGBColor(0x0B, 0x3D, 0x91)
GRIS = RGBColor(0x5D, 0x6B, 0x7E)
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def n(x):
    return f"{x:,}".replace(",", ".") if isinstance(x, int) else str(x).replace(".", ",")


def fecha_larga(iso):
    d = datetime.date.fromisoformat(iso[:10])
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def hipervinculo(parrafo, url, texto):
    rel = parrafo.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F6FEB")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    parrafo._p.append(h)


def sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def tabla(doc, cab, filas, anchos=None, derecha=()):
    t = doc.add_table(rows=1, cols=len(cab))
    t.style = "Table Grid"
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cab):
        cel = t.rows[0].cells[i]
        cel.text = ""
        run = cel.paragraphs[0].add_run(c)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sombrear(cel, "0B3D91")
    for f in filas:
        cels = t.add_row().cells
        for i, v in enumerate(f):
            cels[i].text = ""
            p = cels[i].paragraphs[0]
            run = p.add_run("" if v is None else str(v))
            run.font.size = Pt(9)
            if i in derecha:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if anchos:
        for fila in t.rows:
            for i, a in enumerate(anchos):
                fila.cells[i].width = Cm(a)
    doc.add_paragraph()
    return t


def viñeta(doc, texto, negrita=None):
    p = doc.add_paragraph(style="List Bullet")
    if negrita:
        p.add_run(negrita).bold = True
    p.add_run(texto)
    return p


def construir_doc(m, version, url_portal, repo, entorno):
    ctx, k = m["ctx"], {x["id"]: x for x in m["kpi"]}
    op = [t for t in m["tablas"].values() if t["k"] in "rp" and not t["particion"]]
    oper = [t for t in op if not t["apoyo"]]
    cols = sum(len(t["cols"]) for t in op)
    fecha = fecha_larga(ctx["fecha"])
    enlace = lambda vista: (url_portal + "#" + vista) if url_portal else ""

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.3)
        s.top_margin = s.bottom_margin = Cm(2.0)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)
    for nombre, tam in (("Heading 1", 15), ("Heading 2", 12.5)):
        doc.styles[nombre].font.color.rgb = AZUL
        doc.styles[nombre].font.size = Pt(tam)

    # ---- portada breve
    p = doc.add_paragraph()
    r = p.add_run("Escuela Superior de Administración Pública · SuperApp ESAP")
    r.font.color.rgb = GRIS
    r.font.size = Pt(10)
    titulo = doc.add_paragraph()
    r = titulo.add_run("Resumen ejecutivo de la base de datos")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = AZUL
    p = doc.add_paragraph()
    r = p.add_run(f"Entorno {entorno} · base {ctx['db']} · corte del {fecha} · versión {version}")
    r.font.color.rgb = GRIS
    if url_portal:
        p = doc.add_paragraph("Detalle navegable en el Portal de Datos ESAP: ")
        hipervinculo(p, url_portal, url_portal)

    # ---- 1. propósito
    doc.add_heading("1. Propósito", level=1)
    doc.add_paragraph(
        "Este documento resume el estado de la base de datos de la SuperApp ESAP como parte de su entrega: qué contiene, "
        "cómo está organizada, cómo se relacionan sus datos y qué tan documentada y consistente es su estructura. "
        "Está pensado para una lectura rápida. El detalle de cada tabla, columna y relación está en el Portal de Datos ESAP, "
        "y las consultas, los archivos CSV y la documentación en Markdown que lo respaldan están en el repositorio técnico (sección 7).")

    # ---- 2. fuente y alcance
    doc.add_heading("2. Fuente y alcance", level=1)
    doc.add_paragraph(
        f"Las cifras provienen de una extracción de solo lectura de la base {ctx['db']} del entorno {entorno}, hecha el {fecha}. "
        "La extracción consulta el catálogo de PostgreSQL, las estadísticas del planificador y conteos agregados; "
        "no extrae ni publica valores de datos personales.")
    tabla(doc, ["Elemento", "Valor"], [
        ["Motor", ctx["version"].split(" on ")[0].split(" (")[0]],
        ["Tamaño de la base", f"{ctx['bytes'] / 1048576:,.1f} MB".replace(",", "X").replace(".", ",").replace("X", ".")],
        ["Dominios (esquemas)", n(len(m["esquemas"]))],
        ["Tablas", f"{n(len(op))} ({n(len(op) - len(oper))} de apoyo: respaldos, copias o migraciones)"],
        ["Vistas", n(sum(1 for t in m["tablas"].values() if t["k"] in "vm"))],
        ["Columnas", n(cols)],
        ["Relaciones declaradas (FK)", n(len(m["fks"]))],
        ["Funciones y procedimientos · triggers", f"{n(len(m['programables']))} · {n(len(m['triggers']))}"],
        ["Filas (conteo exacto)", n(sum(t["filas"] or 0 for t in op))],
    ], anchos=[6, 10])

    # ---- 3. dominios
    doc.add_heading("3. Organización por dominios", level=1)
    doc.add_paragraph("Cada esquema de la base corresponde a un dominio funcional de la SuperApp:")
    tabla(doc, ["Dominio", "Esquema", "Tablas", "Columnas", "Filas", "Sin filas"],
          [[e["nombre"], e["s"], n(e["tablas"]), n(e["columnas"]), n(e["filas"]), n(e["vacias"])] for e in m["esquemas"]],
          anchos=[5, 4.2, 1.6, 1.8, 2.2, 1.6], derecha=(2, 3, 4, 5))

    # ---- 4. indicadores
    doc.add_heading("4. Indicadores de calidad del modelo", level=1)
    doc.add_paragraph("Los indicadores se calculan sobre las tablas operativas (sin las de apoyo) y se agrupan por dimensión de calidad "
                      "de datos según ISO/IEC 25012. La meta indica el valor de referencia para una base lista para producción.")
    tabla(doc, ["Id", "Indicador", "Dimensión", "Valor", "Meta"],
          [[x["id"], x["nombre"], x["dim"], n(x["valor"]) + (" %" if "%" in x["meta"] else ""), x["meta"]] for x in m["kpi"]],
          anchos=[1.2, 6.5, 3.6, 2, 2.2], derecha=(3,))

    # ---- 5. hallazgos
    doc.add_heading("5. Hallazgos principales", level=1)
    ident = sorted([i for i in m["identidad"] if not i["es_maestro"] and not i["apoyo"] and i["nn"]], key=lambda i: -i["dist"])
    esquemas_ident = sorted({i["tabla"].split(".")[0] for i in ident})
    maestros = sorted({i["maestro"] for i in m["identidad"] if i["maestro"]})
    if ident:
        ej = "; ".join(f"{i['tabla']}.{i['col']} ({n(i['fuera'])} de {n(i['dist'])} valores fuera de la maestra)" for i in ident[:3])
        viñeta(doc, f" Hay {n(len(ident))} columnas de correo o documento en {n(len(esquemas_ident))} esquemas "
                    f"({', '.join(esquemas_ident)}) que repiten datos de identidad en lugar de referenciar a la persona en la fuente maestra "
                    f"({', '.join(maestros)}). Ejemplos: {ej}. Cuando la misma persona se registra en varios módulos, sus datos pueden divergir.",
               "Identidad replicada.")
    imp_ok = [i for i in m["implicitas"] if i["estado"] in ("confirmada por datos", "parcial")]
    if m["implicitas"]:
        viñeta(doc, f" Además de las {n(len(m['fks']))} FK declaradas, {n(len(m['implicitas']))} columnas apuntan por su nombre a la clave de otra tabla sin una FK que lo garantice; "
                    f"los datos confirman total o parcialmente {n(len(imp_ok))} de ellas. Sin la FK, la base no impide registros que apunten a datos inexistentes "
                    f"(indicador K04: {n(k['K04']['valor'])} % de las relaciones están declaradas).", "Relaciones sin restricción.")
    viñeta(doc, f" {n(k['K01']['valor'])} % de las tablas y {n(k['K02']['valor'])} % de las columnas tienen descripción registrada en la base (meta ≥ 95 %). "
                "Las demás aparecen en el portal como «por documentar».", "Documentación del diccionario.")
    vacias = [t for t in oper if t["filas"] == 0]
    if vacias or len(op) > len(oper):
        viñeta(doc, f" {n(len(vacias))} tablas operativas no tienen filas en {entorno} y {n(len(op) - len(oper))} tablas son de apoyo (respaldos, copias o migraciones). "
                    "Pueden corresponder a módulos aún sin uso en este entorno o a estructuras por depurar.", "Tablas sin uso aparente.")
    sin_pk = [t for t in oper if not t["pk"]]
    if sin_pk:
        viñeta(doc, f" {n(len(sin_pk))} tablas operativas no declaran clave primaria ({', '.join(t['s'] + '.' + t['t'] for t in sin_pk[:6])}{'…' if len(sin_pk) > 6 else ''}).",
               "Tablas sin clave primaria.")
    sin_idx = [f for f in m["fks"] if not f["indexada"]]
    if sin_idx:
        viñeta(doc, f" {n(len(sin_idx))} de {n(len(m['fks']))} FK no tienen índice en la columna hija, lo que puede hacer lentas las consultas que unen esas tablas "
                    "y los borrados en la tabla padre.", "Índices de relaciones.")
    pii = [c for t in oper for c in t["cols"] if c["pii"]]
    if pii:
        viñeta(doc, f" {n(len(pii))} columnas contienen datos personales según su nombre ({n(sum(1 for c in pii if c['pii']['nivel'] == 'alta'))} de nivel alto: "
                    "identificación, credenciales o datos sensibles). Su tratamiento debe responder a la Ley 1581 de 2012.", "Datos personales.")

    # ---- 6. recomendaciones
    doc.add_heading("6. Recomendaciones", level=1)
    recs = [
        ("Fuente única de identidad. ", f"Tomar {', '.join(maestros) or 'la tabla maestra de personas'} como la fuente de los datos de contacto e identificación, "
                                         "y que los módulos guarden solo la referencia a la persona (su identificador), resolviendo correo y documento en el momento de la consulta."),
        ("Declarar las relaciones confirmadas. ", "Convertir en FK las relaciones implícitas que los datos confirman, después de depurar los valores sin registro padre, "
                                                   "y crear índices en las columnas hijas de las FK que no lo tienen."),
        ("Controles de calidad en la entrada. ", "Validar formato y obligatoriedad de los datos antes de guardarlos (rechazo o cuarentena de los registros que no cumplan), "
                                                  "para que las tablas operativas solo reciban datos válidos."),
        ("Depuración preventiva. ", "Revisar con el equipo de desarrollo las tablas sin filas y las de apoyo; marcar las que no se usen (por ejemplo con el sufijo _deprecated) "
                                    "durante un periodo de observación y retirarlas después si ninguna aplicación las consulta."),
        ("Completar el diccionario. ", "Registrar las descripciones de negocio como COMMENT ON TABLE / COMMENT ON COLUMN, de modo que la documentación viaje con la base "
                                       "y el portal se actualice al regenerarlo."),
        ("Gobierno de datos personales. ", "Validar la clasificación de columnas con datos personales con el responsable del tratamiento y definir controles de acceso "
                                           "y enmascaramiento para los entornos que no son de producción."),
    ]
    for i, (a, b) in enumerate(recs, 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(a).bold = True
        p.add_run(b)

    # ---- 7. dónde está el detalle
    doc.add_heading("7. Dónde consultar el detalle", level=1)
    p = doc.add_paragraph("Portal de Datos ESAP " + version + ": ")
    if url_portal:
        hipervinculo(p, url_portal, url_portal)
    else:
        p.add_run(f"portal/Portal_de_Datos_ESAP_{version}.html (archivo HTML autónomo; se abre en cualquier navegador, sin conexión).")
    secciones = [("Catálogo", "catalogo", "buscador de tablas y columnas; ficha de cada tabla con columnas, claves, relaciones, índices y uso"),
                 ("Modelo ER", "modelo", "diagrama entidad-relación por dominio, con FK y relaciones implícitas"),
                 ("Relaciones", "relaciones", "multiplicidad UML medida con datos y cobertura de las relaciones implícitas"),
                 ("Calidad", "calidad", "indicadores K01–K10 y listas de tablas sin clave primaria, sin filas y de apoyo"),
                 ("Identidad", "identidad", "réplicas de correo y documento frente a la fuente maestra"),
                 ("Datos personales", "personales", "columnas clasificadas según la Ley 1581 de 2012")]
    t = tabla(doc, ["Sección del portal", "Qué contiene", "Enlace"],
              [[s, d, "" if url_portal else f"#{v}"] for s, v, d in secciones], anchos=[3.3, 9.7, 3])
    if url_portal:
        for fila, (s, v, d) in zip(t.rows[1:], secciones):
            hipervinculo(fila.cells[2].paragraphs[0], enlace(v), "Abrir sección")
    p = doc.add_paragraph("Repositorio técnico: ")
    if repo:
        hipervinculo(p, repo, repo)
    else:
        p.add_run("(ubicación por definir)")
    for ruta, d in [("docs/csv/", "catálogo, columnas, relaciones, identidad, datos personales e indicadores en CSV"),
                    ("docs/catalogo/", "un documento Markdown por dominio con todas sus tablas y columnas"),
                    (f"docs/informe_extraccion_{version}.md", "cifras de la extracción y cambios frente a la anterior"),
                    ("sql/extraccion/", "consultas de solo lectura que producen todo lo anterior y permiten regenerarlo")]:
        viñeta(doc, f" {d}.", ruta)

    # ---- anexo
    doc.add_heading("Anexo · Términos", level=1)
    for a, b in [("FK (llave foránea): ", "restricción que garantiza que un valor exista en la tabla referenciada."),
                 ("Relación implícita: ", "columna que por su nombre apunta a otra tabla, pero sin FK que lo garantice."),
                 ("Multiplicidad UML: ", "cuántos registros de cada lado participan en una relación (0..1, 1, 0..*, 1..*)."),
                 ("Tabla de apoyo: ", "respaldo, copia, migración o temporal identificada por su nombre."),
                 ("Fuente maestra: ", "tabla de referencia para un dato (por ejemplo, los datos de identidad de una persona).")]:
        viñeta(doc, b, a)
    p = doc.add_paragraph()
    r = p.add_run(f"Generado el {datetime.date.today():%Y-%m-%d} con herramientas/generar_resumen.py a partir de la extracción del {ctx['fecha'][:10]}.")
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRIS
    return doc


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--extraccion", required=True)
    ap.add_argument("--version", default="v1.0")
    ap.add_argument("--entorno", default="QA")
    ap.add_argument("--url-portal", default="", help="URL donde quedará publicado el portal")
    ap.add_argument("--repo", default="", help="URL del repositorio técnico")
    ap.add_argument("--salida")
    a = ap.parse_args()
    try:
        m = M.construir(os.path.abspath(a.extraccion))
    except M.ErrorExtraccion as e:
        sys.exit(f"No se generó el resumen: {e}")
    doc = construir_doc(m, a.version, a.url_portal, a.repo, a.entorno)
    destino = a.salida or os.path.join(M.RAIZ, "informe", f"Resumen_Ejecutivo_BD_ESAP_{a.version}.docx")
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    doc.save(destino)
    print("Resumen ejecutivo:", os.path.relpath(destino, M.RAIZ))


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""Utilidades comunes para los documentos Word del proyecto (resumen ejecutivo e informe interno)."""
import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x0B, 0x3D, 0x91)
GRIS = RGBColor(0x5D, 0x6B, 0x7E)
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def n(x):
    """Número con formato colombiano: 1.234 y 12,5."""
    if isinstance(x, int):
        return f"{x:,}".replace(",", ".")
    if isinstance(x, float):
        return f"{x:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return str(x)


def fecha_larga(iso):
    d = datetime.date.fromisoformat(iso[:10])
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def documento():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.3)
        s.top_margin = s.bottom_margin = Cm(2.0)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for nombre, tam in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        doc.styles[nombre].font.color.rgb = AZUL
        doc.styles[nombre].font.size = Pt(tam)
    return doc


def portada(doc, antetitulo, titulo, subtitulo, marca=None):
    p = doc.add_paragraph()
    r = p.add_run(antetitulo)
    r.font.color.rgb = GRIS
    r.font.size = Pt(10)
    if marca:
        r = p.add_run("   " + marca)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x9A, 0x67, 0x00)
    p = doc.add_paragraph()
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(21)
    r.font.color.rgb = AZUL
    p = doc.add_paragraph()
    r = p.add_run(subtitulo)
    r.font.color.rgb = GRIS


def hipervinculo(parrafo, url, texto):
    rel = parrafo.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F6FEB")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    parrafo._p.append(h)


def _sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def tabla(doc, cab, filas, anchos=None, derecha=()):
    t = doc.add_table(rows=1, cols=len(cab))
    t.style = "Table Grid"
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cab):
        cel = t.rows[0].cells[i]
        cel.text = ""
        run = cel.paragraphs[0].add_run(c)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _sombrear(cel, "0B3D91")
    for f in filas:
        cels = t.add_row().cells
        for i, v in enumerate(f):
            cels[i].text = ""
            p = cels[i].paragraphs[0]
            run = p.add_run("" if v is None else str(v))
            run.font.size = Pt(9)
            if i in derecha:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if anchos:
        for fila in t.rows:
            for i, a in enumerate(anchos):
                fila.cells[i].width = Cm(a)
    doc.add_paragraph()
    return t


def vineta(doc, texto, negrita=None, estilo="List Bullet"):
    p = doc.add_paragraph(style=estilo)
    if negrita:
        p.add_run(negrita).bold = True
    p.add_run(texto)
    return p


def campo(doc, etiqueta, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.add_run(etiqueta + " ").bold = True
    p.add_run(texto)
    return p


def pie(doc, texto):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRIS
